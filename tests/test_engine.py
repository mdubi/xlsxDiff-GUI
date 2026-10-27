"""Regression tests for xlsxDiff engine. Run from project root: python -m unittest discover -s tests -v"""
from __future__ import annotations
from contextlib import contextmanager
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from xlsx_diff_engine import (DiffOptions, ComparisonCancelled, compare_workbooks,
                              parse_index_spec)
from xlsxDiff_GUI import suggest_output


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        d = Path(self.tmp.name)
        self.old = d / 'stara.xlsx'
        self.new = d / 'nowa.xlsx'
        self.out = d / 'DIFF.xlsx'

    def save(self, path, *sheets):
        wb = Workbook()
        wb.remove(wb.active)
        for title, grid in sheets:
            ws = wb.create_sheet(title)
            for row in grid:
                ws.append(row)
        wb.save(path)
        wb.close()

    @contextmanager
    def load(self):
        wb = load_workbook(self.out, read_only=False)
        try:
            yield wb
        finally:
            wb.close()

    def test_changes_colors_sheets_and_literal_formula_safety(self):
        self.save(self.old, ('Wyniki', [['A', 'B'], ['Wynik', 'Tak'], ['Etykieta', '=2+2']]),
                  ('Usunięty', [['Dane']]), ('Bez zmian', [[1, 2]]))
        self.save(self.new, ('Wyniki', [['A', 'B'], ['Wynik', 'Nie'], ['Etykieta', '=3+3']]),
                  ('Dodany', [['Nowe']]), ('Bez zmian', [[1, 2]]))
        stats = compare_workbooks(self.old, self.new, self.out, DiffOptions(compare_formulas=True))
        self.assertEqual((stats.changed_sheets, stats.unchanged_sheets,
                          stats.new_sheets, stats.removed_sheets), (1, 1, 1, 1))
        self.assertEqual(stats.changed_cells, 2)
        with self.load() as wb:
            # Data starting with '=' must not silently become an Excel formula.
            self.assertEqual(wb['Wyniki']['B3'].data_type, 's')
            self.assertIn('=2', wb['Wyniki']['B3'].value)
            self.assertIn('3', wb['Wyniki']['B3'].value)
            self.assertEqual(wb['Dodany'].sheet_properties.tabColor.rgb[-6:].lower(), '0000ff')
            self.assertEqual(wb['Usunięty'].sheet_properties.tabColor.rgb[-6:].lower(), 'ff0000')
            self.assertIsNone(wb['Wyniki'].sheet_properties.tabColor)
            self.assertIsNotNone(wb['Bez zmian'].sheet_properties.tabColor)
        with ZipFile(self.out) as archive:
            richxml = archive.read('xl/sharedStrings.xml').decode('utf-8')
            # Rich fragments written to shared strings; blue/red formatting visible in XML.
            self.assertIn('Nie', richxml)
            self.assertIn('Tak', richxml)

    @staticmethod
    def _rich_runs(path):
        # Extract rich text formatting rather than relying on the flattened
        # openpyxl cell value (which loses the color/strikeout distinctions).
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        with ZipFile(path) as archive:
            root = ET.fromstring(archive.read('xl/sharedStrings.xml'))
        result = []
        for si in root.findall('s:si', ns):
            runs = []
            for run in si.findall('s:r', ns):
                text = run.find('s:t', ns)
                props = run.find('s:rPr', ns)
                runs.append((text.text if text is not None else '',
                             props is not None and props.find('s:strike', ns) is not None,
                             props is not None and props.find('s:u', ns) is not None))
            if runs:
                result.append(runs)
        return result

    def test_whole_numbers_mode_on_excel_numbers_and_percentage_texts(self):
        self.save(self.old, ('Wyniki', [['Wartość', 'Procent', 'Tytuł', 'Liczba separator'],
                                       [0.12345, '12,345%', 'wynik 12,345% jest dobry', '1 234,56'],
                                       [3.5, '-1,23 %', 'Raport A', '1.234,56']]))
        self.save(self.new, ('Wyniki', [['Wartość', 'Procent', 'Tytuł', 'Liczba separator'],
                                       [0.12346, '12,346%', 'wynik 12,346% jest dobry', '1 234,57'],
                                       [3.6, '-1,24 %', 'Raport B', '1.234,57']]))
        options = DiffOptions(whole_number_diff=True)
        stats = compare_workbooks(self.old, self.new, self.out, options)
        self.assertEqual(stats.changed_cells, 8)
        rich_strings = self._rich_runs(self.out)
        for original, revised in [
            ('0.12345', '0.12346'), ('12,345%', '12,346%'),
            ('3.5', '3.6'), ('-1,23 %', '-1,24 %'),
            ('1 234,56', '1 234,57'), ('1.234,56', '1.234,57')
        ]:
            sequences = [runs for runs in rich_strings
                         if [v for v, _strike, _u in runs] == [original, '  →  ', revised]]
            self.assertTrue(sequences, f'No whole-number rich string {original} => {revised}')
            self.assertTrue(sequences[0][0][1], 'Old number must be crossed out')
            self.assertTrue(sequences[0][2][2], 'New number must be underlined')
        # The free-form sentence still uses character-by-character diff.
        self.assertFalse(any([v for v, _, _ in runs] ==
                             ['wynik 12,345% jest dobry', '  →  ', 'wynik 12,346% jest dobry']
                             for runs in rich_strings))

    def test_legacy_numeric_mode_and_date_like_strings(self):
        self.save(self.old, ('Wyniki', [['12,345%', '2026-10-21']]))
        self.save(self.new, ('Wyniki', [['12,346%', '2026-10-22']]))
        compare_workbooks(self.old, self.new, self.out, DiffOptions(whole_number_diff=False))
        self.assertFalse(any([v for v, _, _ in runs] == ['12,345%', '  →  ', '12,346%']
                             for runs in self._rich_runs(self.out)))
        compare_workbooks(self.old, self.new, self.out, DiffOptions(whole_number_diff=True))
        self.assertTrue(any([v for v, _, _ in runs] == ['12,345%', '  →  ', '12,346%']
                            for runs in self._rich_runs(self.out)))
        # Date-looking strings are not treated as numeric values.
        self.assertFalse(any([v for v, _, _ in runs] == ['2026-10-21', '  →  ', '2026-10-22']
                             for runs in self._rich_runs(self.out)))

    def test_cli_opt_in_whole_numbers(self):
        from xlsxDiff import main as cli_main
        from contextlib import redirect_stdout
        import io
        self.save(self.old, ('Wyniki', [['12,345%']]))
        self.save(self.new, ('Wyniki', [['12,346%']]))
        with redirect_stdout(io.StringIO()):
            cli_main([str(self.old), str(self.new), str(self.out), '-N', '-q'])
        self.assertTrue(any([v for v, _, _ in runs] == ['12,345%', '  →  ', '12,346%']
                            for runs in self._rich_runs(self.out)))

    def test_unchanged_cell_and_skip_empty(self):
        self.save(self.old, ('Arkusz', [['text', None]]))
        self.save(self.new, ('Arkusz', [['text', None]]))
        stats = compare_workbooks(self.old, self.new, self.out, DiffOptions(skip_empty=True))
        self.assertEqual(stats.changed_cells, 0)
        self.assertEqual(stats.unchanged_sheets, 1)

    def test_xlsm_input_and_mixed_extensions_output_has_no_vba(self):
        """Macro-enabled extension is allowed as INPUT, output always .xlsx."""
        from zipfile import ZipFile
        old = self.old.with_suffix('.xlsm')
        new = self.new.with_suffix('.xlsm')
        self.save(old, ('Arkusz', [['ID', 'Value'], [1, 42]]))
        self.save(new, ('Arkusz', [['ID', 'Value'], [1, 43]]))
        for a, b in ((old, new), (self.old, new), (old, self.new)):
            if a == self.old:
                self.save(self.old, ('Arkusz', [['ID', 'Value'], [1, 42]]))
            if b == self.new:
                self.save(self.new, ('Arkusz', [['ID', 'Value'], [1, 43]]))
            stats = compare_workbooks(a, b, self.out, DiffOptions(whole_number_diff=True))
            self.assertEqual(stats.changed_cells, 1)
            with self.load() as wb:
                self.assertEqual(wb['Arkusz']['B2'].value, '42  →  43')
            with ZipFile(self.out) as zf:
                self.assertNotIn('xl/vbaProject.bin', zf.namelist())

    def test_xlsm_output_not_allowed(self):
        self.save(self.old, ('Arkusz', [[1]]))
        self.save(self.new, ('Arkusz', [[2]]))
        with self.assertRaisesRegex(ValueError, 'rozszerzenie .xlsx'):
            compare_workbooks(self.old, self.new, self.out.with_suffix('.xlsm'))

    def test_formula_mode(self):
        self.save(self.old, ('Arkusz', [['Formuła'], ['=SUM(1,2)']]))
        self.save(self.new, ('Arkusz', [['Formuła'], ['=SUM(1,3)']]))
        # Formulas have no cached results in newly created openpyxl workbooks.
        data_stats = compare_workbooks(self.old, self.new, self.out)
        self.assertEqual(data_stats.changed_cells, 0)
        formula_stats = compare_workbooks(self.old, self.new, self.out,
                                           DiffOptions(compare_formulas=True))
        self.assertEqual(formula_stats.changed_cells, 1)
        with self.load() as wb:
            self.assertEqual(wb.active['A2'].data_type, 's')
            self.assertIn('SUM(1,', wb.active['A2'].value)

    def test_insert_row_by_index(self):
        self.save(self.old, ('Dane', [['ID', 'Val'], ['A', 10], ['B', 20]]))
        self.save(self.new, ('Dane', [['ID', 'Val'], ['A', 10], ['NEW', 100], ['B', 20]]))
        opts = DiffOptions(index_columns={'Dane': ['A']}, highlight_headers=True, autofilter=True)
        stats = compare_workbooks(self.old, self.new, self.out, opts)
        self.assertEqual(stats.changed_sheets, 1)
        with self.load() as wb:
            ws = wb['Dane']
            self.assertEqual(ws['A3'].value, 'NEW')
            self.assertEqual(ws['B3'].value, '100')
            self.assertEqual(ws['A4'].value, 'B')
            self.assertTrue(ws.auto_filter.ref)

    def test_insert_column_by_header_index(self):
        self.save(self.old, ('Dane', [['A', 'C'], [1, 3]]))
        self.save(self.new, ('Dane', [['A', 'B', 'C'], [1, 2, 3]]))
        opts = DiffOptions(index_rows={'Dane': [1]})
        stats = compare_workbooks(self.old, self.new, self.out, opts)
        self.assertEqual(stats.changed_sheets, 1)
        with self.load() as wb:
            self.assertEqual([wb.active.cell(1, c).value for c in (1, 2, 3)], ['A', 'B', 'C'])

    def test_polish_names_and_special_content(self):
        self.save(self.old, ('Pytania i odpowiedzi', [['Liczba', 'Treść'], [2.3, 'Żółw 🐢']]))
        self.save(self.new, ('Pytania i odpowiedzi', [['Liczba', 'Treść'], [2.4, 'Żubr 🦬']]))
        stats = compare_workbooks(self.old, self.new, self.out)
        self.assertEqual(stats.changed_cells, 2)
        with self.load() as wb:
            self.assertIn('ółw', wb.active['B2'].value)
            self.assertIn('ubr', wb.active['B2'].value)

    def test_output_protection(self):
        self.save(self.old, ('A', [[1]]))
        self.save(self.new, ('A', [[2]]))
        original = self.old.read_bytes()
        with self.assertRaisesRegex(ValueError, 'nadpisać'):
            compare_workbooks(self.old, self.new, self.old)
        self.assertEqual(self.old.read_bytes(), original)

    def test_cancel_does_not_replace_existing_output(self):
        self.save(self.old, ('A', [[1, 2, 3], [4, 5, 6]]))
        self.save(self.new, ('A', [[2, 3, 4], [5, 6, 7]]))
        self.out.write_bytes(b'keep me')
        cancelled = False
        def progress(frac, msg):
            nonlocal cancelled
            if 'kolumna' in msg:
                cancelled = True
        with self.assertRaises(ComparisonCancelled):
            compare_workbooks(self.old, self.new, self.out, progress=progress,
                              is_cancelled=lambda: cancelled)
        self.assertEqual(self.out.read_bytes(), b'keep me')
        self.assertFalse(list(self.out.parent.glob('.xlsxDiff_*.xlsx')))

    def test_indexes_and_paths(self):
        self.assertEqual(parse_index_spec('Dane i liczby!B,C;Drugi!AA', 'column'),
                         {'Dane i liczby': ['B', 'C'], 'Drugi': ['AA']})
        self.assertEqual(parse_index_spec('Dane!1,2', 'row'), {'Dane': [1, 2]})
        for invalid in ('Dane', 'Dane!XYZ!!A', 'Dane!0'):
            with self.assertRaises(ValueError):
                parse_index_spec(invalid, 'row')
        self.assertIn('DIFF_', suggest_output(str(self.old), str(self.new)))

    def test_author_reference_files(self):
        refdir = Path(__file__).resolve().parent
        out = Path(self.tmp.name) / 'ref.xlsx'
        stats = compare_workbooks(refdir / 'in1.xlsx', refdir / 'in2.xlsx', out)
        self.assertGreater(stats.changed_cells, 0)
        self.assertGreater(stats.changed_sheets, 0)
        wb = load_workbook(out, read_only=True)
        try:
            self.assertIn('staff', wb.sheetnames)
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
