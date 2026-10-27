# -*- coding: utf-8 -*-
"""xlsxDiff comparison engine adapted from Rafał Czeczótka's xlsxDiff v2.0.0.

Original source: https://github.com/rafal-dot/xlsxDiff
Copyright (c) 2020-2023 Rafał Czeczótka.
2026 modifications by Michał Dubrawski with ChatGPT (GPT-6) assistance.
Adaptations: module API, incremental progress/cancel, input safety, atomic saving,
input validation, richer diagnostics and several edge-case fixes.

Licensed under the GNU Affero General Public License, version 3 or later.
See LICENSE included with this program.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Callable, Optional
import math
import os
import re
import tempfile
from numbers import Real

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
import xlsxwriter

__version__ = '1.2.1'

Progress = Callable[[float, str], None]
Warning = Callable[[str], None]
CancelCheck = Callable[[], bool]


class ComparisonCancelled(Exception):
    """Comparison aborted before overwriting the destination."""


@dataclass(frozen=True)
class DiffOptions:
    # Equivalent of the original xlsxDiff flags:
    compare_formulas: bool = False      # -f
    highlight_headers: bool = False     # -x
    index_columns: dict[str, list[str]] = field(default_factory=dict)  # -c: row alignment
    index_rows: dict[str, list[int]] = field(default_factory=dict)     # -r: column alignment
    highlight_added_removed: bool = True    # inverse of -X
    autofilter: bool = False            # -a, only when highlight_headers
    skip_empty: bool = False           # -e
    whole_number_diff: bool = False    # --whole-numbers (CLI opt-in; GUI defaults ON)


@dataclass
class DiffStats:
    changed_cells: int = 0
    unchanged_cells: int = 0
    new_sheets: int = 0
    removed_sheets: int = 0
    changed_sheets: int = 0
    unchanged_sheets: int = 0
    warnings: list[str] = field(default_factory=list)


class _Engine:
    def __init__(self, options: DiffOptions, progress: Optional[Progress],
                 is_cancelled: Optional[CancelCheck], warn: Optional[Warning]):
        self.options = options
        self.progress = progress or (lambda frac, message: None)
        self.is_cancelled = is_cancelled or (lambda: False)
        self.warn_callback = warn or (lambda message: None)
        self.stats = DiffStats()
        self._workbook = None
        self._formats = {}

    def check_cancel(self):
        if self.is_cancelled():
            raise ComparisonCancelled('Porównanie anulowano przez użytkownika.')

    def warn(self, msg):
        self.stats.warnings.append(msg)
        self.warn_callback(msg)

    def make_formats(self, wb):
        f = {}
        common = {'text_wrap': True, 'border': 1, 'align': 'top'}
        f['simple'] = wb.add_format(common)
        f['added'] = wb.add_format({**common, 'bold': True, 'underline': True, 'color': 'blue'})
        f['removed'] = wb.add_format({**common, 'bold': True, 'font_strikeout': True, 'color': 'red'})
        f['added_row'] = wb.add_format({**common, 'bold': True, 'underline': True,
                                        'color': 'blue', 'bg_color': '#E8E8FF'})
        f['removed_row'] = wb.add_format({**common, 'bold': True, 'font_strikeout': True,
                                          'color': 'red', 'bg_color': '#FFE8E8'})
        f['modified'] = wb.add_format({**common, 'bg_color': '#A9D171'})
        f['added_modified'] = wb.add_format({**common, 'bold': True, 'underline': True,
                                             'color': 'blue', 'bg_color': '#A9D171'})
        f['removed_modified'] = wb.add_format({**common, 'bold': True, 'font_strikeout': True,
                                               'color': 'red', 'bg_color': '#A9D171'})
        f['equal'] = wb.add_format({**common, 'bg_color': '#C0C0C0'})
        self._formats = f

    def format_for(self, row, col, modified_rows, modified_cols,
                   standard, header=None, update=False, added_or_removed=None):
        f = self._formats
        if update and self.options.highlight_headers:
            modified_rows.add(row)
            modified_cols.add(col)
        if self.options.highlight_headers and (col == 0 and row in modified_rows or
                                               row == 0 and col in modified_cols):
            return f[header or 'modified']
        if added_or_removed and self.options.highlight_added_removed:
            return f[added_or_removed]
        return f[standard]

    @staticmethod
    def _val(ws, row, col):
        """Zero-based location; -1 / out-of-bounds indicates a blank cell."""
        if row < 0 or col < 0 or row >= ws.max_row or col >= ws.max_column:
            return ''
        value = ws.cell(row + 1, col + 1).value
        return '' if value is None else str(value)

    @staticmethod
    def _indexes(a, b):
        """Diff-index alignment, adapted from upstream, preserving order."""
        aligned = []
        for tag, i1, i2, j1, j2 in SequenceMatcher(None, a, b).get_opcodes():
            if tag == 'equal':
                aligned.extend((i, j) for i, j in zip(range(i1, i2), range(j1, j2)))
            elif tag == 'delete':
                aligned.extend((i, -1) for i in range(i1, i2))
            elif tag == 'insert':
                aligned.extend((-1, j) for j in range(j1, j2))
            else:
                aligned.extend((i, -1) for i in range(i1, i2))
                aligned.extend((-1, j) for j in range(j1, j2))
        # Descending order important: top-left headers are written last.
        return list(reversed([(out, i, j) for out, (i, j) in enumerate(aligned)]))

    def row_alignment(self, name, old, new):
        identifiers = self.options.index_columns.get(name)
        if not identifiers:
            return list(reversed([(i, i, i) for i in range(max(old.max_row, new.max_row))]))
        cols = [column_index_from_string(c) - 1 for c in identifiers]
        def row_key(ws, row):
            return tuple(self._val(ws, row, c) for c in cols)
        a = [row_key(old, i) for i in range(old.max_row)]
        b = [row_key(new, i) for i in range(new.max_row)]
        if len(set(a)) < len(a) or len(set(b)) < len(b):
            self.warn(f'{name}: nieunikalne identyfikatory wierszy (-c); dopasowanie może być niejednoznaczne.')
        return self._indexes(a, b)

    def col_alignment(self, name, old, new):
        identifiers = self.options.index_rows.get(name)
        if not identifiers:
            return list(reversed([(i, i, i) for i in range(max(old.max_column, new.max_column))]))
        rows = [r - 1 for r in identifiers]
        def col_key(ws, col):
            return tuple(self._val(ws, r, col) for r in rows)
        a = [col_key(old, i) for i in range(old.max_column)]
        b = [col_key(new, i) for i in range(new.max_column)]
        if len(set(a)) < len(a) or len(set(b)) < len(b):
            self.warn(f'{name}: nieunikalne identyfikatory kolumn (-r); dopasowanie może być niejednoznaczne.')
        return self._indexes(a, b)

    def safe_string(self, ws, row, col, value, fmt):
        # Crucial security/data-correctness fix: never reinterpret Excel-looking
        # values such as '=SUM(...)' or a URL as formulas / hyperlinks.
        if len(value) > 32767:
            self.warn(f'{ws.name}!{get_column_letter(col + 1)}{row + 1}: '
                      'tekst przekracza limit Excela (32767 znaków), został skrócony.')
            value = value[:32752] + ' [SKRÓCONO]'
        if value == '':
            ws.write_blank(row, col, None, fmt)
        else:
            ws.write_string(row, col, value, fmt)

    def write_changed_rich(self, ws, row, col, old_text, new_text, fmt):
        """Cell-level red strikeout deletions / blue underlined insertions."""
        # Avoid XlsxWriter silent -2 truncation of a rich string.
        rich = []
        for tag, i1, i2, j1, j2 in SequenceMatcher(None, old_text, new_text).get_opcodes():
            if tag == 'equal':
                if i1 != i2:
                    rich.append(old_text[i1:i2])
            elif tag == 'delete':
                if i1 != i2:
                    rich.extend([self._formats['removed'], old_text[i1:i2]])
            elif tag == 'insert':
                if j1 != j2:
                    rich.extend([self._formats['added'], new_text[j1:j2]])
            else:
                if i1 != i2:
                    rich.extend([self._formats['removed'], old_text[i1:i2]])
                if j1 != j2:
                    rich.extend([self._formats['added'], new_text[j1:j2]])
        # Rich fragment total may exceed 32767 even when both input values are valid.
        text_size = sum(len(x) for x in rich if isinstance(x, str))
        if text_size > 32767:
            self.warn(f'{ws.name}!{get_column_letter(col + 1)}{row + 1}: '
                      'suma starych i nowych fragmentów przekracza limit Excela; '
                      'pokazano skrócony zapis tekstowy.')
            self.safe_string(ws, row, col,
                             f'PRZED: {old_text[:15000]}\nPO: {new_text[:15000]}', fmt)
            return
        if len(rich) < 3:
            # XlsxWriter requires multiple fragments for a rich string.
            self.safe_string(ws, row, col, f'PRZED: {old_text}\nPO: {new_text}', fmt)
            return
        ret = ws.write_rich_string(row, col, *rich, fmt)
        if ret != 0:
            self.warn(f'{ws.name}!{get_column_letter(col + 1)}{row + 1}: '
                      f'zapis fragmentów nie powiódł się (kod {ret}); użyto tekstu.')
            self.safe_string(ws, row, col, f'PRZED: {old_text}\nPO: {new_text}', fmt)

    # Textual numeric values are common in survey questionnaires: e.g. "12,34 %",
    # "-2.5", "1 234,56". Date-like strings such as "2026-10-21" do not match.
    # This is a recognition step, NOT numeric normalization or rounding: the
    # original cell values are compared and displayed exactly as before.
    _NUMERIC_TEXT = re.compile(
        r"^[+-]?(?:[0-9]+(?:[.,][0-9]+)?|[0-9]{1,3}(?:[ .\u00a0\u202f,][0-9]{3})+(?:[.,][0-9]+)?)"
        r"(?:[eE][+-]?[0-9]+)?\s*%?$"
    )

    @classmethod
    def _is_numeric_cell(cls, cell):
        value = cell.value
        # bool inherits int in Python; dates can be numeric in XLSX internals,
        # but openpyxl represents them as datetime/date objects (excluded).
        if isinstance(value, Real) and not isinstance(value, bool):
            return math.isfinite(value)
        return isinstance(value, str) and bool(cls._NUMERIC_TEXT.fullmatch(value.strip()))

    def write_whole_number_rich(self, ws, row, col, old_text, new_text, fmt):
        """Show a changed number as whole red strikeout and whole blue underline."""
        arrow = '  →  '
        if len(old_text) + len(new_text) + len(arrow) > 32767:
            self.warn(f'{ws.name}!{get_column_letter(col + 1)}{row + 1}: '
                      'dwie liczby razem przekraczają limit Excela; zapis uproszczony.')
            self.safe_string(ws, row, col,
                             f'PRZED: {old_text[:15000]}\nPO: {new_text[:15000]}', fmt)
            return
        result = ws.write_rich_string(
            row, col, self._formats['removed'], old_text, arrow,
            self._formats['added'], new_text, fmt)
        if result != 0:
            self.warn(f'{ws.name}!{get_column_letter(col + 1)}{row + 1}: '
                      f'nie można zapisać pełnych liczb (kod {result}); zapis tekstowy.')
            self.safe_string(ws, row, col, f'PRZED: {old_text}\nPO: {new_text}', fmt)

    def compare_cell(self, out, ro, co, old, r1, c1, new, r2, c2, modified_rows, modified_cols):
        v1 = self._val(old, r1, c1)
        v2 = self._val(new, r2, c2)
        opts = self.options
        if (r1 < 0 and c2 < 0) or (c1 < 0 and r2 < 0):
            out.write_blank(ro, co, None, self.format_for(ro, co, modified_rows, modified_cols, 'equal'))
            return False
        if r1 < 0 or c1 < 0:
            fmt = self.format_for(ro, co, modified_rows, modified_cols, 'added',
                                  'added_modified', True, 'added_row')
            self.safe_string(out, ro, co, v2, fmt)
            if v2:
                self.stats.changed_cells += 1
            return bool(v2)
        if r2 < 0 or c2 < 0:
            fmt = self.format_for(ro, co, modified_rows, modified_cols, 'removed',
                                  'removed_modified', True, 'removed_row')
            self.safe_string(out, ro, co, v1, fmt)
            if v1:
                self.stats.changed_cells += 1
            return bool(v1)
        if v1 == v2:
            self.stats.unchanged_cells += 1
            if v1:
                self.safe_string(out, ro, co, v1,
                                 self.format_for(ro, co, modified_rows, modified_cols,
                                                 'equal', 'modified'))
            elif not opts.skip_empty:
                out.write_blank(ro, co, None,
                                self.format_for(ro, co, modified_rows, modified_cols,
                                                'equal', 'modified'))
            return False
        self.stats.changed_cells += 1
        if not v1 or not v2:
            added = not v1
            fmt = self.format_for(ro, co, modified_rows, modified_cols,
                                  'added' if added else 'removed',
                                  'added_modified' if added else 'removed_modified', True)
            self.safe_string(out, ro, co, v2 if added else v1, fmt)
        else:
            fmt = self.format_for(ro, co, modified_rows, modified_cols,
                                  'simple', 'modified', True)
            if (opts.whole_number_diff
                    and self._is_numeric_cell(old.cell(r1 + 1, c1 + 1))
                    and self._is_numeric_cell(new.cell(r2 + 1, c2 + 1))):
                self.write_whole_number_rich(out, ro, co, v1, v2, fmt)
            else:
                self.write_changed_rich(out, ro, co, v1, v2, fmt)
        return True

    @staticmethod
    def _width(ws, col):
        if col < 0:
            return 2
        width = ws.column_dimensions[get_column_letter(col + 1)].width
        return width if isinstance(width, (float, int)) and math.isfinite(width) else 8.43

    def compare_sheet(self, out, name, old, new, index, total):
        rows = self.row_alignment(name, old, new)
        cols = self.col_alignment(name, old, new)
        if not rows or not cols:
            self.stats.unchanged_sheets += 1
            out.set_tab_color('gray')
            return
        if rows[0][0] >= 1048576 or cols[0][0] >= 16384:
            raise ValueError(f'Arkusz {name}: po uwzględnieniu wstawek przekracza limity XLSX.')
        modified_rows, modified_cols = set(), set()
        changed = False
        for j, (co, c1, c2) in enumerate(cols):
            self.check_cancel()
            out.set_column(co, co, max(self._width(old, c1), self._width(new, c2), 2))
            for k, (ro, r1, r2) in enumerate(rows):
                if not k % 128:
                    self.check_cancel()
                if self.compare_cell(out, ro, co, old, r1, c1, new, r2, c2,
                                     modified_rows, modified_cols):
                    changed = True
            if j == 0 or j % 8 == 0 or j == len(cols) - 1:
                self.progress((index + (j + 1) / len(cols)) / total,
                              f'{name}: kolumna {j + 1}/{len(cols)}')
        if changed:
            self.stats.changed_sheets += 1
            if self.options.autofilter and self.options.highlight_headers:
                out.autofilter(0, 0, 0, cols[0][0])
        else:
            out.set_tab_color('gray')
            self.stats.unchanged_sheets += 1

    def clone_sheet(self, out, name, ws, operation, index, total):
        is_added = operation == 'added'
        fmt = self._formats['added' if is_added else 'removed']
        out.set_tab_color('blue' if is_added else 'red')
        if is_added:
            self.stats.new_sheets += 1
        else:
            self.stats.removed_sheets += 1
        for c in range(ws.max_column):
            self.check_cancel()
            out.set_column(c, c, self._width(ws, c))
            for r in range(ws.max_row):
                if r % 128 == 0:
                    self.check_cancel()
                self.safe_string(out, r, c, self._val(ws, r, c), fmt)
            if c == 0 or c % 8 == 0 or c == ws.max_column - 1:
                self.progress((index + (c + 1) / max(ws.max_column, 1)) / total,
                              f'{name}: kopiowanie kolumny {c + 1}/{ws.max_column}')

    def run(self, old_path: Path, new_path: Path, output_path: Path):
        self.check_cancel()
        self.progress(0, 'Wczytywanie starszej wersji…')
        old_wb = load_workbook(old_path, data_only=not self.options.compare_formulas,
                               read_only=False, keep_vba=False, keep_links=False)
        try:
            self.check_cancel()
            self.progress(0, 'Wczytywanie nowszej wersji…')
            new_wb = load_workbook(new_path, data_only=not self.options.compare_formulas,
                                   read_only=False, keep_vba=False, keep_links=False)
            try:
                self.check_cancel()
                names = list(dict.fromkeys([*new_wb.sheetnames, *old_wb.sheetnames]))
                for name in set(self.options.index_columns) | set(self.options.index_rows):
                    if name not in names:
                        self.warn(f'Nie znaleziono arkusza {name!r} z ustawień indeksów.')
                # Workbook writes a temporary output. Target stays unchanged until success.
                output_path.parent.mkdir(parents=True, exist_ok=True)
                fd, temporary = tempfile.mkstemp(suffix='.xlsx', prefix='.xlsxDiff_', dir=output_path.parent)
                os.close(fd)
                try:
                    self._workbook = xlsxwriter.Workbook(temporary, {
                        'strings_to_formulas': False,
                        'strings_to_urls': False,
                    })
                    self.make_formats(self._workbook)
                    try:
                        for i, name in enumerate(names):
                            self.check_cancel()
                            self.progress(i / len(names), f'Porównywanie arkusza {i+1}/{len(names)}: {name}')
                            out = self._workbook.add_worksheet(name)
                            if name in old_wb and name in new_wb:
                                self.compare_sheet(out, name, old_wb[name], new_wb[name], i, len(names))
                            elif name in new_wb:
                                self.clone_sheet(out, name, new_wb[name], 'added', i, len(names))
                            else:
                                self.clone_sheet(out, name, old_wb[name], 'removed', i, len(names))
                    finally:
                        self._workbook.close()
                        self._workbook = None
                    self.check_cancel()
                    os.replace(temporary, output_path)
                finally:
                    if os.path.exists(temporary):
                        os.unlink(temporary)
            finally:
                new_wb.close()
        finally:
            old_wb.close()
        self.progress(1, 'Gotowe')
        return self.stats


def _validate_options(options: DiffOptions):
    for sheet, cols in options.index_columns.items():
        if not sheet or not cols:
            raise ValueError('Każda pozycja -c musi mieć nazwę arkusza i listę kolumn.')
        for col in cols:
            try:
                number = column_index_from_string(col)
                if number > 16384:
                    raise ValueError('Kolumna poza limitem Excela.')
            except (ValueError, TypeError) as err:
                raise ValueError(f'Nieprawidłowa kolumna indeksowa: {sheet}!{col}') from err
    for sheet, rows in options.index_rows.items():
        if not sheet or not rows or not all(isinstance(r, int) and 1 <= r <= 1048576 for r in rows):
            raise ValueError(f'Nieprawidłowe indeksy wierszy: {sheet}: {rows}')


def compare_workbooks(old_file: str | Path, new_file: str | Path,
                      output_file: str | Path, options: Optional[DiffOptions] = None,
                      progress: Optional[Progress] = None, is_cancelled: Optional[CancelCheck] = None,
                      warn: Optional[Warning] = None) -> DiffStats:
    """Compare workbook contents and generate a Word-style highlighted XLSX report.

    Warning: numeric values are stringified, number formats aren't preserved,
    and data_only=True reads cached formula results, NOT recalculated values.
    """
    options = options or DiffOptions()
    _validate_options(options)
    old, new, output = (Path(p).expanduser().resolve() for p in (old_file, new_file, output_file))
    for file in (old, new):
        if not file.is_file() or file.suffix.lower() not in ('.xlsx', '.xlsm'):
            raise ValueError(f'Plik wejściowy nie istnieje lub nie jest XLSX/XLSM: {file}')
    if old == new:
        raise ValueError('Stary i nowy plik są identyczne.')
    if output == old or output == new:
        raise ValueError('Raport nie może nadpisać pliku źródłowego!')
    if output.suffix.lower() != '.xlsx':
        raise ValueError('Plik raportu musi mieć rozszerzenie .xlsx.')
    if output.name.startswith('~$'):
        raise ValueError('Plik raportu nie może być tymczasowym plikiem blokady Excela.')
    if options.autofilter and not options.highlight_headers:
        raise ValueError('Autofiltr wymaga zaznaczenia podświetlania nagłówków (-x).')
    return _Engine(options, progress, is_cancelled, warn).run(old, new, output)


def parse_index_spec(spec: str, kind: str) -> dict:
    """Parse multiline or semicolon separated Sheet!A,B / Sheet!1,2 specs.

    This parser does NOT split on whitespace, so tab names may include spaces.
    """
    result = {}
    for text in spec.replace(';', '\n').splitlines():
        text = text.strip()
        if not text:
            continue
        if '!' not in text:
            raise ValueError(f'Brakuje ! w identyfikatorze: {text!r} (np. Wyniki!A,B)')
        name, values = text.rsplit('!', 1)
        name = name.strip().strip("'")
        tokens = [x.strip().upper() for x in values.split(',') if x.strip()]
        if not name or not tokens:
            raise ValueError(f'Błędna definicja indeksów: {text!r}')
        if name in result:
            raise ValueError(f'Arkusz {name!r} powtarza się na liście; połącz indeksy w jeden wpis.')
        if kind == 'column':
            for token in tokens:
                try:
                    if column_index_from_string(token) > 16384:
                        raise ValueError('too large')
                except ValueError as err:
                    raise ValueError(f'Błędna kolumna {token!r} w {text!r}') from err
            result[name] = tokens
        elif kind == 'row':
            try:
                numbers = [int(token) for token in tokens]
            except ValueError as err:
                raise ValueError(f'Błędny numer wiersza w {text!r}') from err
            if any(n < 1 or n > 1048576 for n in numbers):
                raise ValueError(f'Wiersz poza zakresem Excela: {text!r}')
            result[name] = numbers
        else:
            raise ValueError(f'Nieobsługiwany rodzaj indeksów: {kind}')
    return result
